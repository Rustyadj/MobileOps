"""Import pipeline: malformed input, validation, duplicate handling, confirmation, no-write and rollback guarantees."""
import io
import zipfile
from datetime import timedelta

import pytest
from openpyxl import Workbook, load_workbook

import imports as data_imports

from .docs_support import DOMAIN, NOW, FaultyDB, env, make_xlsx, server, snapshot  # noqa: F401  (env is a fixture)

pytestmark = pytest.mark.anyio

HEAD = ["Name", "Category", "QR", "Quantity", "Daily Rate", "Condition", "Location"]


async def stage(env, headers, rows, dataset="equipment", **kw):
    return await data_imports.stage_import(
        env.db, dataset=dataset, filename="t.xlsx", content=make_xlsx(headers, rows), actor_id="boss", actor_name="Boss", **kw
    )


async def job_of(env, preview):
    return await env.db.import_jobs.find_one({"id": preview["import_id"]}, {"_id": 0})


async def commit(env, preview, **kw):
    return await data_imports.commit_import(
        env.db, server, import_id=preview["import_id"], plan_hash=preview["plan_hash"], actor_id="boss", actor_name="Boss", **kw
    )


async def audit_events(env):
    return [d["event_type"] for d in await env.db.operational_activity.find({"source": "import"}, {"_id": 0}).to_list(None)]


# ----------------------------- malformed files -----------------------------
def zip_with(content: bytes, name: str, payload: bytes) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content)) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr(name, payload)
    return out.getvalue()


GOOD = make_xlsx(HEAD, [["Router", "tool", "R-1", 2, 5, "good", "Yard"]])


@pytest.mark.parametrize("label,content,message", [
    ("empty", b"", "empty"),
    ("plain text", b"name,qty\nsaw,1\n", "Unsupported file"),
    ("html named xlsx", b"<html><body>nope</body></html>", "Unsupported file"),
    ("zip but not a workbook", b"PK\x03\x04junkjunkjunk", "valid .xlsx"),
    ("truncated workbook", GOOD[: len(GOOD) // 2], "valid .xlsx"),
    ("macro workbook", zip_with(GOOD, "xl/vbaProject.bin", b"\x00evil"), "Macro-enabled"),
    ("zip bomb", zip_with(GOOD, "xl/bomb.bin", b"\x00" * 62_000_000), "unreasonable size"),
    ("oversized upload", b"PK\x03\x04" + b"\x00" * 5_000_001, "too large"),
    ("fake pdf", b"%PDF-1.4\nthis is not really a pdf", "Could not read the PDF"),
])
async def test_malformed_files_are_rejected_without_staging_anything(env, label, content, message):
    before = await snapshot(env.db, *DOMAIN, "import_jobs", "import_files")
    with pytest.raises(data_imports.ImportFileError, match=message):
        await data_imports.stage_import(env.db, dataset="equipment", filename="x", content=content, actor_id="boss", actor_name="Boss")
    assert await snapshot(env.db, *DOMAIN, "import_jobs", "import_files") == before


async def test_too_many_rows_rejected(env):
    rows = [[f"Item {i}", "tool", f"Q{i}", 1, 0, "good", ""] for i in range(data_imports.MAX_ROWS + 1)]
    with pytest.raises(data_imports.ImportFileError, match="Too many rows"):
        await stage(env, HEAD, rows)


async def test_sheet_without_header_row_rejected(env):
    wb = Workbook()
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(data_imports.ImportFileError, match="no header"):
        await data_imports.stage_import(env.db, dataset="equipment", filename="e", content=buf.getvalue(), actor_id="b", actor_name="B")


async def test_unsupported_dataset_and_policy_rejected(env):
    for dataset in ("rentals", "dispatches", "returns", "users"):
        with pytest.raises(data_imports.ImportFileError, match="Imports are supported for"):
            await stage(env, HEAD, [], dataset=dataset)
    with pytest.raises(data_imports.ImportFileError, match="on_duplicate"):
        await stage(env, HEAD, [["x", "tool", "", 1, 0, "good", ""]], on_duplicate="overwrite")


async def test_formula_cells_are_rejected_per_row_not_evaluated(env):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(HEAD)
    ws.append(["Good Row", "tool", "G-1", 1, 0, "good", "Yard"])
    ws.append(["=1+1", "tool", "F-1", 1, 0, "good", "Yard"])
    buf = io.BytesIO()
    wb.save(buf)
    preview = await data_imports.stage_import(env.db, dataset="equipment", filename="f.xlsx", content=buf.getvalue(), actor_id="b", actor_name="B")
    assert preview["summary"]["create"] == 1 and preview["summary"]["error"] == 1
    assert "formulas are not allowed" in preview["errors"][0]["errors"][0]
    assert preview["errors"][0]["row"] == 3


# ----------------------------- mapping & validation ------------------------
async def test_column_aliases_are_detected_and_unknown_columns_reported(env):
    preview = await stage(env, ["Tool", "Type", "Tag", "Qty", "Rate", "Serial #", "Favourite colour"],
                          [["Planer", "tool", "P-1", 3, 7.5, "S9", "blue"]])
    assert preview["mapping"]["Tool"] == "name" and preview["mapping"]["Tag"] == "qr_code" and preview["mapping"]["Qty"] == "quantity"
    assert preview["mapping"]["Serial #"] == "serial_number" and preview["mapping"]["Favourite colour"] is None
    assert preview["ignored_columns"] == ["Favourite colour"]
    row = (await job_of(env, preview))["rows"][0]
    assert row["action"] == "create" and row["data"]["quantity"] == 3 and row["data"]["available"] == 3 and row["data"]["daily_rate"] == 7.5


async def test_missing_required_column_is_a_blocking_error_and_cannot_commit(env):
    preview = await stage(env, ["Qty", "Location"], [[1, "Yard"]])
    assert preview["blocking_errors"] and "name" in preview["blocking_errors"][0]
    assert preview["summary"]["total"] == 0
    with pytest.raises(data_imports.ImportStateError, match="blocking"):
        await commit(env, preview)


async def test_explicit_mapping_overrides_detection(env):
    preview = await stage(env, ["Widget", "Kind", "Count"], [["Clamp", "tool", 4]], mapping={"Widget": "name", "Kind": "category", "Count": "quantity"})
    assert (await job_of(env, preview))["rows"][0]["data"]["name"] == "Clamp"
    with pytest.raises(data_imports.ImportFileError, match="not in the file"):
        await stage(env, HEAD, [], mapping={"Nope": "name"})
    with pytest.raises(data_imports.ImportFileError, match="Unknown target field"):
        await stage(env, HEAD, [], mapping={"Name": "price_of_everything"})
    with pytest.raises(data_imports.ImportFileError, match="More than one column"):
        await stage(env, HEAD, [], mapping={"Name": "name", "Category": "name"})


@pytest.mark.parametrize("row,expected", [
    (["", "tool", "A1", 1, 0, "good", ""], "name: required"),
    (["X", "tool", "A1", "many", 0, "good", ""], "quantity: 'many' is not a number"),
    (["X", "tool", "A1", -3, 0, "good", ""], "quantity: must be at least 0"),
    (["X", "tool", "A1", 1.5, 0, "good", ""], "quantity: must be a whole number"),
    (["X", "tool", "A1", 1, -2, "good", ""], "daily_rate: must be at least 0"),
    (["X", "tool", "A1", 1, "free", "good", ""], "daily_rate: 'free' is not a number"),
    (["X", "tool", "A1", 1, 0, "sparkly", ""], "condition: 'sparkly' is not one of"),
    (["X" * 300, "tool", "A1", 1, 0, "good", ""], "name: longer than 200"),
])
async def test_row_level_validation_reports_the_row_and_reason(env, row, expected):
    preview = await stage(env, HEAD, [["Fine", "tool", "OK-1", 1, 0, "good", ""], row])
    assert preview["summary"]["create"] == 1 and preview["summary"]["error"] == 1
    error = preview["errors"][0]
    assert error["row"] == 3 and any(expected in e for e in error["errors"]) and error["raw"]["Name"].startswith(row[0][:10])


async def test_available_cannot_exceed_quantity(env):
    preview = await stage(env, ["Name", "Category", "Quantity", "Available"], [["Jack", "tool", 2, 5]])
    assert "available: cannot exceed quantity" in preview["errors"][0]["errors"]


async def test_blank_rows_skipped_and_title_rows_above_header_tolerated(env):
    # real-world sheets start with a title; the header is the best-matching row, not blindly the first
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventory 2026"
    ws.append([])
    ws.append(["Yard count"] + [None] * 6)
    ws.append(HEAD)
    ws.append(["A", "tool", "A-1", 1, 0, "good", ""])
    ws.append([None] * 7)
    ws.append(["B", "tool", "B-1", 1, 0, "good", ""])
    buf = io.BytesIO()
    wb.save(buf)
    preview = await data_imports.stage_import(env.db, dataset="equipment", filename="t.xlsx", content=buf.getvalue(), actor_id="b", actor_name="B")
    assert not preview["blocking_errors"] and preview["summary"]["create"] == 2 and preview["summary"]["total"] == 2
    assert [r["row"] for r in (await job_of(env, preview))["rows"]] == [4, 6], "row numbers must match the spreadsheet"


async def test_sheet_with_no_recognisable_header_is_blocked_not_guessed(env):
    preview = await stage(env, ["Foo", "Bar"], [["a", "b"]])
    assert preview["blocking_errors"] and preview["summary"]["total"] == 0


async def test_tools_dataset_forces_category_and_block_forces_kind(env):
    preview = await stage(env, ["Name", "Category", "QR"], [["Wrench", "tool", "W-1"], ["Scaffold", "scaffold", "W-2"]], dataset="tools")
    assert preview["summary"]["create"] == 1 and preview["summary"]["error"] == 1
    assert "this import only accepts 'tool'" in preview["errors"][0]["errors"][0]
    preview = await stage(env, ["Product", "On Hand"], [["Brick", 5]], dataset="block")
    assert (await job_of(env, preview))["rows"][0]["data"]["kind"] == "block"


# ----------------------------- duplicates ----------------------------------
async def test_duplicates_within_the_file_are_errors_naming_the_first_row(env):
    preview = await stage(env, HEAD, [["Router", "tool", "R-9", 1, 0, "good", ""], ["Router again", "tool", "r-9", 1, 0, "good", ""]])
    assert preview["summary"]["create"] == 1 and preview["summary"]["error"] == 1
    assert preview["errors"][0]["errors"] == ["duplicate of row 2 in this file"]


async def test_existing_records_are_skipped_by_default_with_the_would_be_diff(env):
    before = await snapshot(env.db, *DOMAIN)
    preview = await stage(env, HEAD, [["Circular Saw Pro", "tool", "Q-101", 9, 99.0, "good", "Annex"]])
    assert preview["summary"]["skip_duplicate"] == 1 and preview["summary"]["create"] == 0
    row = preview["rows"][0]
    assert row["existing_name"] == "Circular Saw" and row["diff"]["name"] == ["Circular Saw", "Circular Saw Pro"] and row["diff"]["daily_rate"] == [8.0, 99.0]
    assert {"quantity", "location"} <= set(row["ignored_changes"]), "stock/location are ledger-driven and never applied"
    result = await commit(env, preview)
    assert result["result"]["created"] == 0 and result["result"]["updated"] == 0 and result["result"]["skipped_duplicates"] == 1
    assert await snapshot(env.db, *DOMAIN) == before, "a skipped duplicate must leave production data untouched"


async def test_update_mode_applies_only_safe_fields_and_shows_old_and_new(env):
    preview = await stage(env, HEAD, [["Circular Saw Pro", "tool", "Q-101", 9, 99.0, "fair", "Annex"]], on_duplicate="update")
    assert preview["summary"]["update"] == 1
    assert set(preview["rows"][0]["diff"]) == {"name", "daily_rate", "condition"}
    await commit(env, preview)
    saw = await env.db.equipment.find_one({"qr_code": "Q-101"}, {"_id": 0})
    assert (saw["name"], saw["daily_rate"], saw["condition"]) == ("Circular Saw Pro", 99.0, "fair")
    assert (saw["quantity"], saw["available"], saw["location"]) == (2, 2, "Yard"), "stock buckets and location must not change"
    assert await env.db.ledger_entries.count_documents({}) == 0


async def test_ambiguous_match_is_an_error_not_a_guess(env):
    # QR belongs to the Hammer Drill but the SKU belongs to the Circular Saw
    preview = await stage(env, ["Name", "Category", "QR", "SKU"], [["Mixed", "tool", "Q-100", "QR-Q-101"]])
    assert preview["summary"]["error"] == 1 and "ambiguous" in preview["errors"][0]["errors"][0]


async def test_name_only_rows_match_existing_by_name_and_model(env):
    preview = await stage(env, ["Name", "Category"], [["strongback 8ft", "strongback"]])
    assert preview["summary"]["skip_duplicate"] == 1


async def test_sellable_duplicates_by_sku_and_update_diff(env):
    preview = await stage(env, ["Product", "SKU", "On Hand", "Price"], [["Foam Adhesive", "FA-1", 20, 11.0], ["Fresh Item", "NEW-1", 3, 1.0]],
                          dataset="consumables", on_duplicate="update")
    assert preview["summary"]["update"] == 1 and preview["summary"]["create"] == 1
    diff = next(r for r in preview["rows"] if r["action"] == "update")["diff"]
    assert diff == {"quantity_on_hand": [12, 20], "price": [9.0, 11.0]}
    await commit(env, preview)
    assert (await env.db.sellable_items.find_one({"sku": "FA-1"}))["quantity_on_hand"] == 20
    assert (await env.db.sellable_items.find_one({"sku": "NEW-1"}))["kind"] == "consumable"


# ----------------------------- staging & confirmation ----------------------
async def test_staging_and_review_never_write_domain_data(env):
    before = await snapshot(env.db, *DOMAIN)
    preview = await stage(env, HEAD, [["New One", "tool", "N-1", 5, 1, "good", "Yard"], ["Circular Saw", "tool", "Q-101", 2, 8, "good", "Yard"]])
    data_imports.preview(await job_of(env, preview))
    assert await snapshot(env.db, *DOMAIN) == before
    assert preview["summary"]["create"] == 1 and "Nothing has been changed" in preview["next_step"]
    assert "import_staged" in await audit_events(env)


async def test_commit_requires_the_reviewed_plan_hash(env):
    preview = await stage(env, HEAD, [["New One", "tool", "N-1", 5, 1, "good", "Yard"]])
    before = await snapshot(env.db, *DOMAIN)
    for bad in ("", "0" * 64, None):
        with pytest.raises(data_imports.ImportStateError, match="plan_hash"):
            await data_imports.commit_import(env.db, server, import_id=preview["import_id"], plan_hash=bad, actor_id="boss", actor_name="Boss")
    assert await snapshot(env.db, *DOMAIN) == before
    assert (await job_of(env, preview))["status"] == "staged"


async def test_invalid_rows_block_commit_unless_explicitly_skipped(env):
    preview = await stage(env, HEAD, [["Good", "tool", "G-1", 1, 0, "good", ""], ["", "tool", "B-1", 1, 0, "good", ""]])
    before = await snapshot(env.db, *DOMAIN)
    with pytest.raises(data_imports.ImportStateError, match="invalid"):
        await commit(env, preview)
    assert await snapshot(env.db, *DOMAIN) == before
    result = await commit(env, preview, skip_invalid_rows=True)
    assert result["result"]["created"] == 1 and result["result"]["skipped_invalid_rows"] == 1
    assert await env.db.equipment.count_documents({"qr_code": "G-1"}) == 1 and await env.db.equipment.count_documents({"qr_code": "B-1"}) == 0


async def test_commit_creates_equipment_with_ledger_entries_and_audit(env):
    preview = await stage(env, HEAD, [["Planer", "tool", "P-1", 3, 4.5, "good", "Yard"], ["Empty Shelf", "tool", "P-2", 0, 0, "good", "Yard"]])
    result = await commit(env, preview)
    assert result["result"]["created"] == 2
    planer = await env.db.equipment.find_one({"qr_code": "P-1"}, {"_id": 0})
    assert planer["available"] == 3 and planer["location_balances"] == {"Yard": 3} and planer["sku"] == "QR-P-1"
    entries = await env.db.ledger_entries.find({"equipment_id": planer["id"]}, {"_id": 0}).to_list(None)
    assert len(entries) == 1 and entries[0]["qty"] == 3 and entries[0]["reason"] == "received" and "import" in entries[0]["note"]
    assert await env.db.ledger_entries.count_documents({}) == 1, "a zero-stock item gets no ledger entry"
    events = await audit_events(env)
    assert events.count("import_staged") == 1 and events.count("import_committed") == 1
    committed = await env.db.operational_activity.find_one({"event_type": "import_committed"})
    assert committed["requesting_user_name"] == "Boss" and committed["result"]["created"] == 2


async def test_an_import_cannot_be_committed_twice(env):
    preview = await stage(env, HEAD, [["Planer", "tool", "P-1", 3, 4.5, "good", "Yard"]])
    await commit(env, preview)
    count = await env.db.equipment.count_documents({})
    with pytest.raises(data_imports.ImportStateError, match="already committed"):
        await commit(env, preview)
    assert await env.db.equipment.count_documents({}) == count


async def test_concurrent_commit_in_progress_is_refused(env):
    preview = await stage(env, HEAD, [["Planer", "tool", "P-1", 3, 4.5, "good", "Yard"]])
    await env.db.import_jobs.update_one({"id": preview["import_id"]}, {"$set": {"status": "committing"}})
    with pytest.raises(data_imports.ImportStateError, match="already committing"):
        await commit(env, preview)
    assert await env.db.equipment.count_documents({"qr_code": "P-1"}) == 0


async def test_expired_preview_cannot_be_committed(env):
    preview = await stage(env, HEAD, [["Planer", "tool", "P-1", 3, 4.5, "good", "Yard"]])
    await env.db.import_jobs.update_one({"id": preview["import_id"]}, {"$set": {"expires_at": NOW - timedelta(days=400)}})
    with pytest.raises(data_imports.ImportStateError, match="expired"):
        await commit(env, preview)
    assert await env.db.equipment.count_documents({"qr_code": "P-1"}) == 0


async def test_commit_refuses_when_data_changed_since_preview(env):
    preview = await stage(env, HEAD, [["Planer", "tool", "P-1", 3, 4.5, "good", "Yard"]])
    # someone adds the same QR through the app after the admin looked at the preview
    await env.db.equipment.insert_one(server.Equipment(sku="QR-P-1", qr_code="P-1", name="Planer (manual)", category="tool").model_dump())
    before = await snapshot(env.db, *DOMAIN)
    with pytest.raises(data_imports.ImportStateError, match="changed since"):
        await commit(env, preview)
    assert await snapshot(env.db, *DOMAIN) == before
    assert (await job_of(env, preview))["status"] == "stale"
    assert "import_refused_stale" in await audit_events(env)


async def test_cancel_discards_without_touching_data(env):
    preview = await stage(env, HEAD, [["Planer", "tool", "P-1", 3, 4.5, "good", "Yard"]])
    before = await snapshot(env.db, *DOMAIN)
    out = await data_imports.cancel_import(env.db, preview["import_id"], actor_id="boss", actor_name="Boss")
    assert out["status"] == "cancelled" and await snapshot(env.db, *DOMAIN) == before
    with pytest.raises(data_imports.ImportStateError):
        await commit(env, preview)


# ----------------------------- rollback ------------------------------------
async def test_failure_midway_rolls_back_every_write(env):
    rows = [["Alpha", "tool", "RB-1", 2, 1, "good", "Yard"], ["Circular Saw Pro", "tool", "Q-101", 2, 50, "good", "Yard"],
            ["Beta", "tool", "RB-2", 4, 1, "good", "Yard"], ["Gamma", "tool", "RB-3", 1, 1, "good", "Yard"]]
    preview = await stage(env, HEAD, rows, on_duplicate="update")
    assert preview["summary"]["create"] == 3 and preview["summary"]["update"] == 1
    before = await snapshot(env.db, *DOMAIN)

    async def explode_on_fourth_write(index, row):
        if index == 3:
            raise RuntimeError("disk on fire")

    with pytest.raises(data_imports.ImportCommitError, match="rolled back") as err:
        await commit(env, preview, before_apply=explode_on_fourth_write)
    assert err.value.rolled_back is True
    assert await snapshot(env.db, *DOMAIN) == before, "inserts, the in-place update and ledger entries must all be undone"
    job = await job_of(env, preview)
    assert job["status"] == "failed" and "disk on fire" in job["commit_error"]
    failed = await env.db.operational_activity.find_one({"event_type": "import_failed"})
    assert failed["result"]["rolled_back"] is True and failed["result"]["net_writes"] == 0
    with pytest.raises(data_imports.ImportStateError):
        await commit(env, preview)  # a failed import is terminal; it must be re-previewed


async def test_real_database_error_during_insert_also_rolls_back(env):
    preview = await stage(env, HEAD, [["Alpha", "tool", "RB-1", 2, 1, "good", "Yard"], ["Beta", "tool", "RB-2", 4, 1, "good", "Yard"]])
    before = await snapshot(env.db, *DOMAIN)
    faulty = FaultyDB(env.db, {("equipment", "insert_one"): lambda n: n == 2})  # second equipment insert dies
    with pytest.raises(data_imports.ImportCommitError, match="rolled back"):
        await data_imports.commit_import(faulty, server, import_id=preview["import_id"], plan_hash=preview["plan_hash"], actor_id="boss", actor_name="Boss")
    assert faulty.calls[("equipment", "insert_one")] == 2
    assert await snapshot(env.db, *DOMAIN) == before, "the first item and its ledger entry must be gone"


async def test_error_while_writing_the_ledger_rolls_back_the_item_too(env):
    preview = await stage(env, HEAD, [["Alpha", "tool", "RB-1", 2, 1, "good", "Yard"]])
    before = await snapshot(env.db, *DOMAIN)
    faulty = FaultyDB(env.db, {("ledger_entries", "insert_one"): lambda n: True})
    with pytest.raises(data_imports.ImportCommitError):
        await data_imports.commit_import(faulty, server, import_id=preview["import_id"], plan_hash=preview["plan_hash"], actor_id="boss", actor_name="Boss")
    assert await snapshot(env.db, *DOMAIN) == before


async def test_rollback_failure_is_reported_loudly(env):
    preview = await stage(env, HEAD, [["Alpha", "tool", "RB-1", 2, 1, "good", "Yard"], ["Beta", "tool", "RB-2", 4, 1, "good", "Yard"]])
    faulty = FaultyDB(env.db, {("equipment", "insert_one"): lambda n: n == 2, ("equipment", "delete_one"): lambda n: True})
    with pytest.raises(data_imports.ImportCommitError, match="rollback was incomplete") as err:
        await data_imports.commit_import(faulty, server, import_id=preview["import_id"], plan_hash=preview["plan_hash"], actor_id="boss", actor_name="Boss")
    assert err.value.rolled_back is False
    job = await job_of(env, preview)
    assert job["status"] == "failed_needs_review" and job["rollback_errors"]
    failed = await env.db.operational_activity.find_one({"event_type": "import_failed"})
    assert failed["result"]["rolled_back"] is False and failed["result"]["net_writes"] == "unknown"


async def test_plan_report_lists_every_row_with_problems_and_changes(env):
    preview = await stage(env, HEAD, [["Circular Saw Pro", "tool", "Q-101", 9, 99.0, "good", "Annex"], ["", "tool", "X", 1, 0, "good", ""]])
    wb = load_workbook(io.BytesIO(data_imports.render_plan_report(await job_of(env, preview))))
    rows = [[c.value for c in r] for r in wb["Import plan"].iter_rows()]
    assert rows[0][:2] == ["Row", "Action"] and {r[1] for r in rows[1:]} == {"skip_duplicate", "error"}
    assert any("name: required" in (r[3] or "") for r in rows[1:]) and any("daily_rate: 8.0 -> 99.0" in (r[4] or "") for r in rows[1:])
