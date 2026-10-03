"""MCP tools Nathan uses for exports, PDFs and imports: scopes, admin-only commit, audit trail, end-to-end flow."""
import base64
import io

import pytest
from mcp.server.fastmcp.exceptions import ToolError
from openpyxl import load_workbook

from mcp_server import (
    DEFAULT_HERMES_SCOPES, HERMES_AGENT_ID, AgentPrincipal, create_mobileops_mcp, issue_admin_grant,
)

from .docs_support import DOMAIN, env, make_xlsx, server, snapshot  # noqa: F401  (env is a fixture)

pytestmark = pytest.mark.anyio

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEAD = ["Name", "Category", "QR", "Quantity", "Daily Rate", "Condition", "Location"]
PUBLIC = "https://icfops.srv1427612.hstgr.cloud"


def integration(scopes=DEFAULT_HERMES_SCOPES):
    principal = AgentPrincipal(HERMES_AGENT_ID, frozenset(scopes))
    return create_mobileops_mcp(server, principal_provider=lambda: principal)


async def call(mcp, tool, **args):
    return await mcp.mcp._tool_manager.call_tool(tool, args)


def path_of(url: str) -> str:
    assert url.startswith(PUBLIC), url
    return url[len(PUBLIC):]


async def grant(env, admin_id="boss", name="Boss") -> str:
    return await issue_admin_grant(env.db, admin_id=admin_id, admin_name=name, source_message_id="msg-1")


async def audit(env, tool):
    return [d for d in await env.db.mcp_audit_log.find({"tool": tool}, {"_id": 0}).to_list(None)]


def workbook_rows(content: bytes):
    return [[c.value for c in row] for row in load_workbook(io.BytesIO(content))["Data"].iter_rows()]


NEW_ROWS = [["Planer", "tool", "P-1", 3, 4.5, "good", "Yard"], ["Jointer", "tool", "P-2", 1, 6, "good", "Yard"]]


# ----------------------------- registry ------------------------------------
async def test_tools_are_registered_with_honest_annotations():
    tools = integration().mcp._tool_manager._tools
    expected = {"export_report", "rental_agreement_pdf", "dispatch_ticket_pdf", "import_create_upload", "import_preview", "import_status",
                "import_plan_report", "import_commit", "import_cancel"}
    assert expected <= set(tools)
    assert tools["import_commit"].annotations.destructiveHint is True and tools["import_commit"].annotations.readOnlyHint is False
    for name in expected - {"import_commit"}:
        assert tools[name].annotations.readOnlyHint is True, f"{name} writes nothing to MobileOps data, so it must not demand a confirmation"


# ----------------------------- exports -------------------------------------
async def test_export_nicks_tools_to_excel(env):
    out = await call(integration(), "export_report", dataset="assignments", format="xlsx", assigned_to="Nick")
    assert out["ok"] and "confirmation_required" not in out
    data = out["data"]
    assert data["download_url"].startswith(f"{PUBLIC}/api/files/shared/") and data["record_count"] == 1 and data["filters"] == {"checked_out_to": "Nick"}
    assert data["expires_in_seconds"] == 600 and data["filename"].endswith(".xlsx")
    download = await env.client.get(path_of(data["download_url"]))
    assert download.status_code == 200 and download.content[:2] == b"PK"
    rows = workbook_rows(download.content)
    assert rows[0][0] == "Tool" and [r[0] for r in rows[1:]] == ["Hammer Drill"] and "Nick Smith" in rows[1]
    assert data["sha256"]
    log = (await audit(env, "export_report"))[-1]
    assert log["status"] == "succeeded" and log["agent_identity"] == HERMES_AGENT_ID
    activity = await env.db.operational_activity.find_one({"source": "export", "event_type": "xlsx_export"})
    assert activity["parameters"]["channel"] == "mcp" and activity["parameters"]["dataset"] == "assignments"


async def test_export_report_validates_inputs_and_scope(env):
    mcp = integration()
    for args in ({"dataset": "users"}, {"dataset": "equipment", "status": "active"}, {"dataset": "rentals", "format": "docx"},
                 {"dataset": "equipment", "filters": {"$where": "1"}}):
        with pytest.raises(ToolError):
            await call(mcp, "export_report", **args)
    denied = integration(set(DEFAULT_HERMES_SCOPES) - {"operations:read"})
    with pytest.raises(ToolError, match="Missing required scope"):
        await call(denied, "export_report", dataset="equipment", format="xlsx")
    assert (await audit(env, "export_report"))[-1]["status"] == "failed"


async def test_generic_filters_object_and_new_formats(env):
    out = await call(integration(), "export_report", dataset="shop_tasks", format="pdf", filters={"priority": "high"})
    assert out["data"]["record_count"] == 1 and out["data"]["format"] == "pdf"
    assert (await env.client.get(path_of(out["data"]["download_url"]))).content.startswith(b"%PDF")


async def test_rental_and_dispatch_pdfs_over_mcp(env):
    mcp = integration()
    rental = (await call(mcp, "rental_agreement_pdf", rental_id="r1-aaaaaaaa"))["data"]
    assert (await env.client.get(path_of(rental["download_url"]))).content.startswith(b"%PDF") and rental["dataset"] == "rental_agreement"
    ticket = (await call(mcp, "dispatch_ticket_pdf", dispatch_id="d2-cccccccc"))["data"]
    assert ticket["dataset"] == "dispatch_inbound"
    with pytest.raises(ToolError, match="not found"):
        await call(mcp, "rental_agreement_pdf", rental_id="nope")
    with pytest.raises(ToolError, match="Missing required scope"):
        await call(integration(set(DEFAULT_HERMES_SCOPES) - {"rentals:read"}), "rental_agreement_pdf", rental_id="r1-aaaaaaaa")


# ----------------------------- import: staging -----------------------------
async def test_upload_link_flow_stages_a_preview_without_writing(env):
    mcp = integration()
    out = (await call(mcp, "import_create_upload", dataset="equipment"))["data"]
    assert out["single_use"] and out["expires_in_seconds"] == 600 and "changes no MobileOps data" in out["note"]
    before = await snapshot(env.db, *DOMAIN)
    r = await env.client.post(path_of(out["upload_url"]), files={"file": ("inv.xlsx", make_xlsx(HEAD, NEW_ROWS), XLSX_MIME)})
    assert r.status_code == 200 and r.json()["summary"]["create"] == 2
    assert await snapshot(env.db, *DOMAIN) == before


async def test_inline_preview_validates_and_keeps_the_payload_out_of_the_audit_log(env):
    mcp = integration()
    payload = base64.b64encode(make_xlsx(HEAD, NEW_ROWS)).decode()
    preview = (await call(mcp, "import_preview", dataset="equipment", filename="inv.xlsx", file_base64=payload))["data"]
    assert preview["summary"]["create"] == 2 and preview["blocking_errors"] == []
    stored = (await audit(env, "import_preview"))[-1]
    assert payload not in str(stored["parameters"]) and stored["parameters"]["file_base64_chars"] == len(payload)
    with pytest.raises(ToolError, match="not valid base64"):
        await call(mcp, "import_preview", dataset="equipment", filename="x.xlsx", file_base64="@@@@")
    with pytest.raises(ToolError, match="1 MB"):
        await call(mcp, "import_preview", dataset="equipment", filename="x.xlsx", file_base64="A" * 1_500_000)
    with pytest.raises(ToolError, match="Imports are supported"):
        await call(mcp, "import_preview", dataset="rentals", filename="x.xlsx", file_base64=payload)
    page = (await call(mcp, "import_status", import_id=preview["import_id"], action="create", limit=1))["data"]
    assert page["rows_total_matching"] == 2 and len(page["rows"]) == 1


async def test_dataset_scopes_are_enforced(env):
    payload = base64.b64encode(make_xlsx(["Product", "On Hand"], [["Brick", 5]])).decode()
    equipment_only = integration(set(DEFAULT_HERMES_SCOPES) - {"inventory:write"})
    with pytest.raises(ToolError, match="inventory:write"):
        await call(equipment_only, "import_preview", dataset="block", filename="b.xlsx", file_base64=payload)
    with pytest.raises(ToolError, match="inventory:write"):
        await call(equipment_only, "import_create_upload", dataset="consumables")
    inventory_only = integration(set(DEFAULT_HERMES_SCOPES) - {"equipment:write"})
    assert (await call(inventory_only, "import_preview", dataset="block", filename="b.xlsx", file_base64=payload))["ok"]
    with pytest.raises(ToolError, match="equipment:write"):
        await call(inventory_only, "import_create_upload", dataset="tools")
    read_only = integration({s for s in DEFAULT_HERMES_SCOPES if s.endswith(":read")})
    with pytest.raises(ToolError, match="Missing required scope"):
        await call(read_only, "import_create_upload", dataset="equipment")


async def test_plan_report_download_link(env):
    mcp = integration()
    payload = base64.b64encode(make_xlsx(HEAD, NEW_ROWS + [["", "tool", "BAD", 1, 0, "good", ""]])).decode()
    preview = (await call(mcp, "import_preview", dataset="equipment", filename="inv.xlsx", file_base64=payload))["data"]
    report = (await call(mcp, "import_plan_report", import_id=preview["import_id"]))["data"]
    content = (await env.client.get(path_of(report["download_url"]))).content
    rows = [[c.value for c in r] for r in load_workbook(io.BytesIO(content))["Import plan"].iter_rows()]
    assert len(rows) == 4 and any("name: required" in (r[3] or "") for r in rows[1:])


# ----------------------------- import: commit is admin-only ----------------------
async def staged(env, mcp=None, rows=NEW_ROWS, dataset="equipment", headers=HEAD, **kw):
    payload = base64.b64encode(make_xlsx(headers, rows)).decode()
    return (await call(mcp or integration(), "import_preview", dataset=dataset, filename="inv.xlsx", file_base64=payload, **kw))["data"]


async def test_commit_without_an_admin_grant_is_refused_even_with_a_confirmation_token(env):
    mcp = integration()
    preview = await staged(env, mcp)
    before = await snapshot(env.db, *DOMAIN)
    with pytest.raises(ToolError, match="admin-only"):
        await call(mcp, "import_commit", import_id=preview["import_id"], plan_hash=preview["plan_hash"])
    with pytest.raises(ToolError, match="admin-only"):
        await call(mcp, "import_commit", import_id=preview["import_id"], plan_hash=preview["plan_hash"], confirmation_token="anything")
    assert await snapshot(env.db, *DOMAIN) == before
    assert (await audit(env, "import_commit"))[-1]["status"] == "failed"


async def test_commit_with_admin_grant_applies_the_reviewed_plan_and_is_audited(env):
    mcp = integration()
    preview = await staged(env, mcp)
    token = await grant(env)
    out = await call(mcp, "import_commit", import_id=preview["import_id"], plan_hash=preview["plan_hash"], confirmation_token=token)
    assert out["ok"] and out["data"]["result"]["created"] == 2
    assert await env.db.equipment.count_documents({"qr_code": {"$in": ["P-1", "P-2"]}}) == 2
    assert await env.db.ledger_entries.count_documents({}) == 2
    log = (await audit(env, "import_commit"))[-1]
    assert log["status"] == "succeeded" and log["authorized_by"]["admin_id"] == "boss" and log["authorized_by"]["type"] == "admin_grant"
    assert token not in str(log), "the raw grant must never be stored"
    committed = await env.db.operational_activity.find_one({"event_type": "import_committed", "source": "import"})
    assert committed["requesting_user"] == "boss" and committed["requesting_user_name"] == "Nathan (for Boss)"
    # and the same grant cannot re-apply it
    with pytest.raises(ToolError, match="already committed"):
        await call(mcp, "import_commit", import_id=preview["import_id"], plan_hash=preview["plan_hash"], confirmation_token=token)
    assert await env.db.equipment.count_documents({"qr_code": {"$in": ["P-1", "P-2"]}}) == 2


async def test_commit_with_wrong_hash_or_stale_data_writes_nothing(env):
    mcp = integration()
    preview = await staged(env, mcp)
    token = await grant(env)
    before = await snapshot(env.db, *DOMAIN)
    with pytest.raises(ToolError, match="plan_hash"):
        await call(mcp, "import_commit", import_id=preview["import_id"], plan_hash="0" * 64, confirmation_token=token)
    await env.db.equipment.insert_one(server.Equipment(sku="QR-P-1", qr_code="P-1", name="Manual", category="tool").model_dump())
    after_manual = await snapshot(env.db, *DOMAIN)
    with pytest.raises(ToolError, match="changed since"):
        await call(mcp, "import_commit", import_id=preview["import_id"], plan_hash=preview["plan_hash"], confirmation_token=token)
    assert await snapshot(env.db, *DOMAIN) == after_manual and before != after_manual
    assert await env.db.equipment.count_documents({"qr_code": "P-2"}) == 0


async def test_grant_of_a_demoted_admin_and_missing_dataset_scope_are_refused(env):
    mcp = integration()
    preview = await staged(env, mcp)
    stale = await grant(env, admin_id="fore1", name="Foreman")  # grant issued for someone who is not an admin
    with pytest.raises(ToolError, match="no longer backed by an active admin"):
        await call(mcp, "import_commit", import_id=preview["import_id"], plan_hash=preview["plan_hash"], confirmation_token=stale)
    block = await staged(env, mcp, rows=[["Brick", 5]], dataset="block", headers=["Product", "On Hand"])
    no_inventory = integration(set(DEFAULT_HERMES_SCOPES) - {"inventory:write"})
    with pytest.raises(ToolError, match="inventory:write"):
        await call(no_inventory, "import_commit", import_id=block["import_id"], plan_hash=block["plan_hash"], confirmation_token=await grant(env))
    assert await env.db.sellable_items.count_documents({"product": "Brick"}) == 0


async def test_update_import_is_shown_then_applied_only_for_safe_fields(env):
    mcp = integration()
    preview = await staged(env, mcp, rows=[["Circular Saw XL", "tool", "Q-101", 99, 20.0, "good", "Annex"]], on_duplicate="update")
    assert preview["summary"]["update"] == 1
    row = preview["rows"][0]
    assert row["diff"]["name"] == ["Circular Saw", "Circular Saw XL"] and "quantity" in row["ignored_changes"]
    await call(mcp, "import_commit", import_id=preview["import_id"], plan_hash=preview["plan_hash"], confirmation_token=await grant(env))
    saw = await env.db.equipment.find_one({"qr_code": "Q-101"}, {"_id": 0})
    assert (saw["name"], saw["daily_rate"], saw["quantity"], saw["location"]) == ("Circular Saw XL", 20.0, 2, "Yard")


async def test_cancel_over_mcp_discards_the_preview(env):
    mcp = integration()
    preview = await staged(env, mcp)
    out = await call(mcp, "import_cancel", import_id=preview["import_id"])
    assert out["data"]["status"] == "cancelled"
    with pytest.raises(ToolError, match="cancelled"):
        await call(mcp, "import_commit", import_id=preview["import_id"], plan_hash=preview["plan_hash"], confirmation_token=await grant(env))


async def test_conversational_round_trip_export_edit_reimport(env):
    """'Export Nick's tools', the admin edits the sheet, 'import this spreadsheet' -> reviewed update, then applied."""
    mcp = integration()
    exported = (await call(mcp, "export_report", dataset="tools", format="xlsx", filters={"checked_out_to": "nick"}))["data"]
    content = (await env.client.get(path_of(exported["download_url"]))).content
    wb = load_workbook(io.BytesIO(content))
    ws = wb["Data"]
    header = [c.value for c in ws[1]]
    ws.cell(2, header.index("Daily Rate") + 1).value = 15.0
    ws.cell(2, header.index("Notes") + 1).value = "Blue case now"
    edited = io.BytesIO()
    wb.save(edited)
    payload = base64.b64encode(edited.getvalue()).decode()
    preview = (await call(mcp, "import_preview", dataset="tools", filename="nicks-tools.xlsx", file_base64=payload, on_duplicate="update"))["data"]
    assert preview["summary"]["update"] == 1 and preview["summary"]["error"] == 0
    assert set(preview["rows"][0]["diff"]) == {"daily_rate", "notes"}
    drill = await env.db.equipment.find_one({"qr_code": "Q-100"}, {"_id": 0})
    assert drill["daily_rate"] == 12.5, "nothing changes before the admin confirms"
    await call(mcp, "import_commit", import_id=preview["import_id"], plan_hash=preview["plan_hash"], confirmation_token=await grant(env))
    drill = await env.db.equipment.find_one({"qr_code": "Q-100"}, {"_id": 0})
    assert (drill["daily_rate"], drill["notes"], drill["checked_out_to"], drill["checked_out"]) == (15.0, "Blue case now", "Nick Smith", 1)
