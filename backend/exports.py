"""CSV / XLSX / PDF export of MobileOps records.

Pure of Mongo and FastAPI: the API layer loads documents (already scoped by the
caller's role) and this module renders them. Also holds the short-lived signed
download token used when Hermes hands a user a link to a generated file.
"""
from __future__ import annotations

import base64
import csv
import hashlib
import hmac
import io
import json
import re
import time
from dataclasses import dataclass, field
from datetime import datetime, timezone
from typing import Any, Callable, Iterable, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

EXPORT_FORMATS = ("pdf", "csv", "xlsx")
EXPORT_TOKEN_TTL_SECONDS = 600


def _fmt_dt(value: Any) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    return "" if value is None else str(value)


def _lines(doc: dict) -> int:
    return len(doc.get("lines") or doc.get("items") or [])


@dataclass(frozen=True)
class Column:
    key: str
    header: str
    get: Optional[Callable[[dict], Any]] = None
    money: bool = False  # omitted entirely for crew, matching redact_money_for_crew

    def value(self, doc: dict) -> Any:
        raw = self.get(doc) if self.get else doc.get(self.key)
        if isinstance(raw, datetime):
            return _fmt_dt(raw)
        if raw is None:
            return ""
        if self.money:
            return f"{float(raw):.2f}"
        return raw

    def raw_value(self, doc: dict) -> Any:
        """Typed cell value for XLSX: numbers stay numbers, datetimes stay dates."""
        raw = self.get(doc) if self.get else doc.get(self.key)
        if isinstance(raw, datetime):
            return raw.astimezone(timezone.utc).replace(tzinfo=None) if raw.tzinfo else raw
        if raw is None:
            return None
        if isinstance(raw, bool):
            return "Yes" if raw else "No"
        if isinstance(raw, (int, float)):
            return raw
        if isinstance(raw, (list, dict)):
            return str(raw)
        return raw


@dataclass(frozen=True)
class Dataset:
    name: str
    title: str
    columns: tuple[Column, ...]
    filters: tuple[str, ...]  # equality filters a caller may apply
    sort: tuple[str, int]
    collection: str = ""  # Mongo collection; defaults to the dataset name
    base_query: dict = field(default_factory=dict)  # always applied, before caller filters
    ci_filters: tuple[str, ...] = ()  # case-insensitive word-start match (people, places)
    xlsx_extra: tuple[Column, ...] = ()  # appended to the XLSX only, so a re-import keeps full fidelity

    @property
    def source(self) -> str:
        return self.collection or self.name

    def query(self, filters: Optional[dict[str, str]]) -> dict[str, Any]:
        """Mongo query for already-cleaned filters. Values are escaped, never treated as operators."""
        clauses: list[dict[str, Any]] = [dict(self.base_query)] if self.base_query else []
        for key, value in (filters or {}).items():
            if key in self.ci_filters:
                # Word-start match: "Nick" finds "Nick Smith" but not "Dominick".
                clauses.append({key: {"$regex": f"(^|[^A-Za-z0-9]){re.escape(value.strip())}", "$options": "i"}})
            else:
                clauses.append({key: value})
        if not clauses:
            return {}
        return clauses[0] if len(clauses) == 1 else {"$and": clauses}

    def xlsx_columns(self, role: str) -> list[Column]:
        return visible_columns(self, role) + [c for c in self.xlsx_extra if not (c.money and role == "crew")]


def _days_out(doc: dict) -> Any:
    since = doc.get("checked_out_at")
    if not isinstance(since, datetime):
        return ""
    if since.tzinfo is None:
        since = since.replace(tzinfo=timezone.utc)
    return max((datetime.now(timezone.utc) - since).days, 0)


# Importable identity/descriptive fields that the on-screen/PDF table omits for width.
_EQUIPMENT_FULL: tuple[Column, ...] = (
    Column("model", "Model"), Column("equipment_family", "Family"), Column("serial_number", "Serial #"),
    Column("sku", "SKU"), Column("tracking_type", "Tracking"), Column("notes", "Notes"),
)

_CUSTODY: tuple[Column, ...] = (
    Column("name", "Tool"), Column("qr_code", "QR"), Column("category", "Category"),
    Column("checked_out_to", "Assigned To"), Column("checked_out_crew", "Crew"), Column("checked_out_job", "Job"),
    Column("checked_out", "Qty Out"), Column("checked_out_at", "Since"), Column("days_out", "Days Out", get=_days_out),
    Column("expected_return_at", "Due Back"), Column("location", "Location"),
)

_EQUIPMENT_SORT = ("name", 1)

_EQUIPMENT_COLUMNS: tuple[Column, ...] = (
    Column("name", "Name"), Column("qr_code", "QR"), Column("category", "Category"),
    Column("condition", "Condition"), Column("location", "Location"),
    Column("quantity", "Owned"), Column("available", "Available"), Column("reserved", "Reserved"),
    Column("on_rental", "On Rental"), Column("in_maintenance", "In Repair"), Column("missing", "Missing"),
    Column("daily_rate", "Daily Rate", money=True),
)

_DAMAGED_COLUMNS: tuple[Column, ...] = (
    Column("name", "Name"), Column("qr_code", "QR"), Column("category", "Category"),
    Column("condition", "Condition"), Column("in_maintenance", "In Repair"), Column("location", "Location"),
    Column("notes", "Notes"),
)

_SELLABLE_COLUMNS: tuple[Column, ...] = (
    Column("product", "Product"), Column("manufacturer", "Manufacturer"), Column("sku", "SKU"),
    Column("unit", "Unit"), Column("core_size", "Core Size"), Column("form_type", "Form Type"),
    Column("quantity_on_hand", "On Hand"), Column("quantity_reserved", "Reserved"),
    Column("reorder_point", "Reorder Point"), Column("cost", "Cost", money=True),
    Column("price", "Price", money=True), Column("notes", "Notes"),
)

DATASETS: dict[str, Dataset] = {
    "equipment": Dataset(
        "equipment", "Equipment Inventory", _EQUIPMENT_COLUMNS,
        ("category", "condition", "location", "checked_out_to"), ("name", 1),
        ci_filters=("location", "checked_out_to"),
        xlsx_extra=_EQUIPMENT_FULL,
    ),
    "rentals": Dataset(
        "rentals", "Rentals",
        (
            Column("customer_name", "Customer"), Column("job_site", "Job Site"), Column("job_address", "Address"),
            Column("status", "Status"), Column("start_date", "Start"), Column("due_date", "Due"),
            Column("lines", "Lines", get=_lines), Column("deposit", "Deposit", money=True),
            Column("id", "ID", get=lambda d: str(d.get("id", ""))[:8]),
        ),
        ("status",), ("created_at", -1),
    ),
    "maintenance": Dataset(
        "maintenance", "Repair Tickets",
        (
            Column("equipment_name", "Equipment"), Column("issue", "Issue"), Column("status", "Status"),
            Column("assigned_to", "Assigned To"), Column("location", "Location"), Column("qty", "Qty"),
            Column("reported_at", "Reported"), Column("cost", "Cost", money=True),
            Column("id", "Ticket", get=lambda d: str(d.get("id", ""))[:8]),
        ),
        ("status",), ("created_at", -1),
    ),
    "dispatches": Dataset(
        "dispatches", "Dispatches",
        (
            Column("scheduled_date", "Scheduled"), Column("direction", "Direction"), Column("status", "Status"),
            Column("customer_name", "Customer"), Column("job_site", "Job Site"), Column("driver_name", "Driver"),
            Column("truck", "Truck"), Column("lines", "Lines", get=_lines),
            Column("id", "ID", get=lambda d: str(d.get("id", ""))[:8]),
        ),
        ("status", "direction"), ("scheduled_date", 1),
    ),
    "tools": Dataset(
        "tools", "Tools", _EQUIPMENT_COLUMNS,
        ("condition", "location", "checked_out_to"), _EQUIPMENT_SORT,
        collection="equipment", base_query={"category": "tool"},
        ci_filters=("location", "checked_out_to"), xlsx_extra=_EQUIPMENT_FULL,
    ),
    "assignments": Dataset(
        "assignments", "Tool Assignments (Checked Out)", _CUSTODY,
        ("checked_out_to", "checked_out_crew", "checked_out_job", "category"), ("checked_out_to", 1),
        collection="equipment", base_query={"checked_out": {"$gt": 0}},
        ci_filters=("checked_out_to", "checked_out_crew", "checked_out_job"),
    ),
    "damaged": Dataset(
        "damaged", "Damaged Equipment", _DAMAGED_COLUMNS,
        ("category", "location"), _EQUIPMENT_SORT,
        collection="equipment",
        base_query={"$or": [{"in_maintenance": {"$gt": 0}}, {"condition": {"$in": ["poor", "broken", "damaged"]}}]},
        ci_filters=("location",),
    ),
    "returns": Dataset(
        "returns", "Inbound / Returns", (
            Column("scheduled_date", "Scheduled"), Column("status", "Status"), Column("customer_name", "Customer"),
            Column("job_site", "Job Site"), Column("driver_name", "Driver"), Column("truck", "Truck"),
            Column("lines", "Lines", get=_lines), Column("id", "ID", get=lambda d: str(d.get("id", ""))[:8]),
        ),
        ("status",), ("scheduled_date", 1),
        collection="dispatches", base_query={"direction": "inbound"},
    ),
    "outbound": Dataset(
        "outbound", "Outbound Deliveries", (
            Column("scheduled_date", "Scheduled"), Column("status", "Status"), Column("customer_name", "Customer"),
            Column("job_site", "Job Site"), Column("driver_name", "Driver"), Column("truck", "Truck"),
            Column("lines", "Lines", get=_lines), Column("id", "ID", get=lambda d: str(d.get("id", ""))[:8]),
        ),
        ("status",), ("scheduled_date", 1),
        collection="dispatches", base_query={"direction": "outbound"},
    ),
    "shop_tasks": Dataset(
        "shop_tasks", "Shop Tasks", (
            Column("title", "Task"), Column("task_type", "Type"), Column("status", "Status"), Column("priority", "Priority"),
            Column("assignee", "Assignee"), Column("due_date", "Due"), Column("qty", "Qty"),
            Column("completed_at", "Completed"), Column("notes", "Notes"),
            Column("id", "ID", get=lambda d: str(d.get("id", ""))[:8]),
        ),
        ("status", "task_type", "priority", "assignee"), ("created_at", -1),
        ci_filters=("assignee",),
    ),
    "consumables": Dataset(
        "consumables", "Consumables", _SELLABLE_COLUMNS,
        ("manufacturer", "unit"), ("product", 1),
        collection="sellable_items", base_query={"kind": "consumable"}, ci_filters=("manufacturer",),
    ),
    "block": Dataset(
        "block", "Block", _SELLABLE_COLUMNS,
        ("manufacturer", "core_size", "form_type"), ("product", 1),
        collection="sellable_items", base_query={"kind": "block"}, ci_filters=("manufacturer",),
    ),
}

# Natural-language names an assistant is likely to use, mapped to a dataset.
DATASET_ALIASES = {
    "inventory": "equipment", "tool": "tools", "assignment": "assignments", "tool_assignments": "assignments",
    "damaged_equipment": "damaged", "damaged_items": "damaged", "return": "returns", "inbound": "returns",
    "shop": "shop_tasks", "shop_task": "shop_tasks", "tasks": "shop_tasks", "consumable": "consumables",
    "blocks": "block", "repairs": "maintenance",
}


class ExportError(ValueError):
    pass


def resolve_dataset(name: str) -> Dataset:
    key = (name or "").strip().lower().replace("-", "_").replace(" ", "_")
    key = DATASET_ALIASES.get(key, key)
    try:
        return DATASETS[key]
    except KeyError:
        raise ExportError(f"Unknown export '{name}'. Available: {', '.join(sorted(DATASETS))}") from None


def clean_filters(dataset: Dataset, filters: Optional[dict[str, Any]]) -> dict[str, str]:
    """Keep only whitelisted, non-empty equality filters (never raw Mongo operators)."""
    out: dict[str, str] = {}
    for key, value in (filters or {}).items():
        if key not in dataset.filters:
            raise ExportError(f"Unsupported filter '{key}' for {dataset.name}")
        if value not in (None, ""):
            if not isinstance(value, str):
                raise ExportError(f"Filter '{key}' must be text")
            out[key] = value
    return out


def visible_columns(dataset: Dataset, role: str) -> list[Column]:
    return [c for c in dataset.columns if not (c.money and role == "crew")]


def render_csv(dataset: Dataset, docs: Iterable[dict], role: str) -> str:
    columns = visible_columns(dataset, role)
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([c.header for c in columns])
    for doc in docs:
        # Guard against spreadsheet formula injection from user-entered text.
        writer.writerow([_csv_safe(c.value(doc)) for c in columns])
    return buf.getvalue()


XLSX_MAX_ROWS = 50_000
_HEADER_FILL = PatternFill("solid", fgColor="1F2937")
_BAND_FILL = PatternFill("solid", fgColor="F3F4F6")
_THIN = Side(style="thin", color="D1D5DB")


def render_xlsx(
    dataset: Dataset, docs: list[dict], role: str, *, brand: str, generated_by: str,
    filters: Optional[dict[str, str]] = None, generated_at: Optional[datetime] = None,
) -> bytes:
    """Formatted workbook: a `Data` sheet (typed cells, frozen header, filter) and a `Report Info` sheet.

    Text is always stored as text — a value like "=SUM(A1)" typed into a name field must
    stay inert when someone opens the file, and must come back identical on re-import.
    """
    if len(docs) > XLSX_MAX_ROWS:
        raise ExportError(f"Too many rows for one workbook ({len(docs)} > {XLSX_MAX_ROWS}); add a filter")
    columns = dataset.xlsx_columns(role)
    generated_at = generated_at or datetime.now(timezone.utc)
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append([c.header for c in columns])
    for col_idx, col in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=_THIN)
    widths = [len(c.header) for c in columns]
    for row_idx, doc in enumerate(docs, start=2):
        for col_idx, col in enumerate(columns, start=1):
            value = col.raw_value(doc)
            cell = ws.cell(row=row_idx, column=col_idx)
            if isinstance(value, str):
                cell.value = value
                cell.data_type = "s"  # never let openpyxl turn "=..." into a formula
            else:
                cell.value = value
            if isinstance(value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm"
            elif col.money:
                cell.number_format = "#,##0.00"
            elif isinstance(value, float):
                cell.number_format = "0.00"
            if row_idx % 2 == 1:
                cell.fill = _BAND_FILL
            widths[col_idx - 1] = max(widths[col_idx - 1], min(len(str(value if value is not None else "")), 48))
    for col_idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width + 3
    ws.freeze_panes = "A2"
    if docs:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(docs) + 1}"
    ws.sheet_view.zoomScale = 100

    info = wb.create_sheet("Report Info")
    for row in (
        ("Report", f"{brand} — {dataset.title}"),
        ("Dataset", dataset.name),
        ("Generated (UTC)", generated_at.strftime("%Y-%m-%d %H:%M")),
        ("Generated by", generated_by),
        ("Filters", ", ".join(f"{k}={v}" for k, v in (filters or {}).items()) or "none"),
        ("Records", len(docs)),
    ):
        info.append(list(row))
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 60
    for row in info.iter_rows(min_col=1, max_col=1):
        row[0].font = Font(bold=True)
    wb.properties.creator = generated_by
    wb.properties.title = f"{brand} {dataset.title}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _csv_safe(value: Any) -> Any:
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


class _NumberedCanvas(rl_canvas.Canvas):
    """Defers page-number drawing until the page count is known ("Page 2 of 5")."""

    def __init__(self, *args: Any, footer: str = "", **kwargs: Any) -> None:
        super().__init__(*args, **kwargs)
        self._saved: list[dict] = []
        self._footer = footer

    def showPage(self) -> None:
        self._saved.append(dict(self.__dict__))
        self._startPage()

    def save(self) -> None:
        total = len(self._saved)
        for state in self._saved:
            self.__dict__.update(state)
            self.setFont("Helvetica", 8)
            self.setFillColor(colors.HexColor("#666666"))
            width, _ = self._pagesize
            self.drawString(0.5 * inch, 0.35 * inch, self._footer)
            self.drawRightString(width - 0.5 * inch, 0.35 * inch, f"Page {self._pageNumber} of {total}")
            super().showPage()
        super().save()


def render_pdf(
    dataset: Dataset, docs: list[dict], role: str, *, brand: str, generated_by: str,
    filters: Optional[dict[str, str]] = None, generated_at: Optional[datetime] = None,
) -> bytes:
    columns = visible_columns(dataset, role)
    generated_at = generated_at or datetime.utcnow()
    styles = getSampleStyleSheet()
    cell = ParagraphStyle("cell", parent=styles["BodyText"], fontSize=7.5, leading=9)
    head = ParagraphStyle("head", parent=cell, textColor=colors.white, fontName="Helvetica-Bold")
    story: list[Any] = [
        Paragraph(f"{_esc(brand)} — {_esc(dataset.title)}", styles["Title"]),
        Paragraph(
            f"Generated {generated_at.strftime('%Y-%m-%d %H:%M')} UTC by {_esc(generated_by)} "
            f"&nbsp;|&nbsp; Filters: {_esc(', '.join(f'{k}={v}' for k, v in (filters or {}).items()) or 'none')}",
            styles["Normal"],
        ),
        Spacer(1, 10),
    ]
    data = [[Paragraph(_esc(c.header), head) for c in columns]]
    for doc in docs:
        data.append([Paragraph(_esc(str(c.value(doc))), cell) for c in columns])
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f4f6")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d1d5db")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story += [table, Spacer(1, 8), Paragraph(f"<b>Total records: {len(docs)}</b>", styles["Normal"])]

    buf = io.BytesIO()
    footer = f"{brand} · {dataset.title}"
    SimpleDocTemplate(
        buf, pagesize=landscape(letter), leftMargin=0.5 * inch, rightMargin=0.5 * inch,
        topMargin=0.5 * inch, bottomMargin=0.6 * inch, title=f"{brand} {dataset.title}", author=generated_by,
    ).build(story, canvasmaker=lambda *a, **k: _NumberedCanvas(*a, footer=footer, **k))
    return buf.getvalue()


def _esc(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


# ----------------------------- signed download links -----------------------
def _b64(raw: bytes) -> str:
    return base64.urlsafe_b64encode(raw).decode().rstrip("=")


def _unb64(text: str) -> bytes:
    return base64.urlsafe_b64decode(text + "=" * (-len(text) % 4))


def _sign(secret: bytes, body: str) -> str:
    return _b64(hmac.new(secret, b"mobileops-export:" + body.encode(), hashlib.sha256).digest())


def sign_export_token(
    secret: bytes, *, dataset: str, fmt: str, filters: dict[str, str], role: str,
    issued_for: str, now: Optional[float] = None, ttl: int = EXPORT_TOKEN_TTL_SECONDS,
) -> str:
    payload = {
        "ds": dataset, "fmt": fmt, "f": filters, "role": role, "by": issued_for,
        "exp": int((now if now is not None else time.time()) + ttl),
    }
    body = _b64(json.dumps(payload, sort_keys=True, separators=(",", ":")).encode())
    return f"{body}.{_sign(secret, body)}"


def verify_export_token(secret: bytes, token: str, now: Optional[float] = None) -> dict[str, Any]:
    try:
        body, signature = token.split(".", 1)
        if not hmac.compare_digest(signature, _sign(secret, body)):
            raise ExportError("Invalid export link")
        payload = json.loads(_unb64(body))
    except (ValueError, TypeError) as exc:
        raise ExportError("Invalid export link") from exc
    if payload.get("exp", 0) < (now if now is not None else time.time()):
        raise ExportError("Export link expired")
    if payload.get("fmt") not in EXPORT_FORMATS:
        raise ExportError("Invalid export link")
    return payload
